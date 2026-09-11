"""Turn a mail (body + attachments) into one flat blob of text.

Nothing here tries to understand structure.  It flattens everything to text
and lets the Neo ID regex do the finding -- a shortlist is equally findable
whether it arrived as a one-column sheet, a pasted block, or an HTML table.

The contract that matters: anything we could NOT read lands in `warnings`,
and a warning must reach the user.  A silent parse failure on the one mail
that mattered is the only way this project actually hurts.
"""
import csv
import io
import os
from dataclasses import dataclass, field

from bs4 import BeautifulSoup
from openpyxl import load_workbook

SHEET_EXTS = {".xlsx", ".xlsm", ".xltx"}
TEXT_EXTS = {".csv", ".tsv", ".txt"}
# Silently ignoring these would be a correctness bug, so name them loudly.
UNSUPPORTED_EXTS = {".pdf", ".xls", ".doc", ".docx", ".zip", ".png", ".jpg", ".jpeg"}


@dataclass
class Extracted:
    text: str = ""
    sources: list = field(default_factory=list)
    warnings: list = field(default_factory=list)


def from_xlsx(data):
    """Every cell of every sheet, flattened.

    Deliberately header-agnostic: the real CDC file's header was "Neo ID "
    with a trailing space, and the next one may differ again.
    """
    lines = []
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None]
                if cells:
                    lines.append(" ".join(cells))
    finally:
        wb.close()
    return "\n".join(lines)


def from_delimited(data):
    text = _decode(data)
    dialect = "excel-tab" if "\t" in text.split("\n", 1)[0] else "excel"
    rows = csv.reader(io.StringIO(text), dialect=dialect)
    return "\n".join(" ".join(c.strip() for c in row if c) for row in rows)


def from_html(html):
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return soup.get_text("\n")


def _decode(data):
    for enc in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract(body_text=None, body_html=None, attachments=()):
    """Flatten a mail into text.

    `attachments` is an iterable of (filename, bytes).
    """
    out = Extracted()
    chunks = []

    # Both parts, not one or the other. The plain-text alternative sometimes
    # flattens an HTML table badly enough to lose rows; IDs are deduplicated
    # downstream, so overlap between the two costs nothing.
    if body_text:
        chunks.append(body_text)
        out.sources.append("body (text)")
    if body_html:
        chunks.append(from_html(body_html))
        out.sources.append("body (html)")

    for filename, data in attachments:
        ext = os.path.splitext(filename)[1].lower()
        try:
            if ext in SHEET_EXTS:
                chunks.append(from_xlsx(data))
                out.sources.append(filename)
            elif ext in TEXT_EXTS:
                chunks.append(from_delimited(data))
                out.sources.append(filename)
            elif ext in UNSUPPORTED_EXTS:
                out.warnings.append(
                    "{}: {} attachments are not parsed -- open this one by hand".format(
                        filename, ext
                    )
                )
            else:
                out.warnings.append("{}: unknown attachment type".format(filename))
        except Exception as exc:  # a corrupt sheet must not kill the whole mail
            out.warnings.append("{}: failed to parse ({})".format(filename, exc))

    out.text = "\n".join(chunks)
    return out
